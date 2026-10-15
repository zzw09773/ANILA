"""Offline document formats, independent fallback, and source citations."""

from __future__ import annotations

import io
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest

from anila_core.ingestion.chunking_plugins import get_chunker
from anila_core.ingestion.errors import ParseError, RemoteParseError
from anila_core.ingestion.parser_registry import ParserRegistry, ParsedDocument
from anila_core.ingestion.parsers import extract_text


@pytest.fixture(autouse=True)
def no_docling(monkeypatch):
    monkeypatch.setattr(ParserRegistry, "_get_docling_parser", classmethod(lambda cls: None))


def _docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("差旅規定", 1)
    doc.add_paragraph("第十二條：應於出差結束後辦理核銷。")
    table = doc.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [["項目", "金額"], ["交通費", "依實報支"]]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _pptx() -> bytes:
    p = "http://schemas.openxmlformats.org/presentationml/2006/main"
    a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "ppt/presentation.xml",
            f'<p:presentation xmlns:p="{p}" xmlns:r="{r}"><p:sldIdLst><p:sldId id="256" r:id="r2"/><p:sldId id="257" r:id="r1"/></p:sldIdLst></p:presentation>',
        )
        archive.writestr(
            "ppt/_rels/presentation.xml.rels",
            '<Relationships><Relationship Id="r1" Target="slides/slide1.xml"/><Relationship Id="r2" Target="slides/slide2.xml"/></Relationships>',
        )
        archive.writestr(
            "ppt/slides/slide1.xml",
            f'<p:sld xmlns:p="{p}" xmlns:a="{a}"><p:cSld><p:spTree/></p:cSld></p:sld>',
        )
        archive.writestr(
            "ppt/slides/slide2.xml",
            f'<p:sld xmlns:p="{p}" xmlns:a="{a}"><p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>投影片標題</a:t></a:r></a:p></p:txBody></p:sp><p:graphicFrame><a:graphic><a:graphicData><a:tbl><a:tr><a:tc><a:txBody><a:p><a:r><a:t>欄位</a:t></a:r></a:p></a:txBody></a:tc></a:tr><a:tr><a:tc><a:txBody><a:p><a:r><a:t>數值</a:t></a:r></a:p></a:txBody></a:tc></a:tr></a:tbl></a:graphicData></a:graphic></p:graphicFrame></p:spTree></p:cSld></p:sld>',
        )
    return buffer.getvalue()


def test_required_extensions_registered():
    assert {
        ".pdf",
        ".txt",
        ".doc",
        ".docx",
        ".rtf",
        ".md",
        ".ppt",
        ".pptx",
        ".py",
        ".c",
        ".cpp",
    } <= set(ParserRegistry.supported_extensions())


@pytest.mark.parametrize(
    "strategy", ["hierarchical", "fixed", "markdown-aware", "pdf-page", "cjk-sentence", "semantic"]
)
@pytest.mark.parametrize(
    "extension,body",
    [
        (
            "py",
            'import os\n\n@cache\ndef lookup(value):\n    """函式說明"""\n    return value\n\nasync def second():\n    return 2\n',
        ),
        (
            "c",
            '#include <stdio.h>\nint lookup(int value) {\n  const char *s = "}";\n  /* { */ return value;\n}\nint second() { return 2; }\n',
        ),
        (
            "cpp",
            "namespace hospital {\nint lookup() { return 1; }\nint second() { return 2; }\n}\n",
        ),
    ],
)
def test_source_functions_and_line_numbers_are_intact(strategy, extension, body):
    text, metadata, _ = extract_text("來源." + extension, body.encode())
    chunks = get_chunker(strategy).chunk(
        text, metadata, {"size": 1, "overlap": 0, "max_leaf_tokens": 1}
    )
    lines = body.splitlines(keepends=True)
    assert "".join(c.content for c in chunks) == body
    for chunk in chunks:
        assert chunk.content == "".join(
            lines[chunk.metadata["line_start"] - 1 : chunk.metadata["line_end"]]
        )
        assert chunk.metadata["source_filename"] == "來源." + extension
    if extension != "cpp":
        assert any("return value" in c.content and "lookup" in c.content for c in chunks)
    else:
        function_chunks = [c for c in chunks if c.metadata["symbol"] != "module"]
        assert [c.metadata["symbol"] for c in function_chunks] == ["lookup", "second"]


@pytest.mark.parametrize(
    "strategy", ["hierarchical", "fixed", "markdown-aware", "pdf-page", "cjk-sentence", "semantic"]
)
def test_slides_and_tables_survive_tiny_chunk_budget(strategy):
    text, metadata, _ = extract_text("差旅.pptx", _pptx())
    chunks = get_chunker(strategy).chunk(
        text, metadata, {"size": 1, "overlap": 0, "max_leaf_tokens": 1}
    )
    assert metadata["slide_count"] == 2
    assert len(chunks) == 1  # second physical slide is empty; it retains its slot
    assert chunks[0].metadata["slide"] == 1
    assert chunks[0].metadata["total_pages"] == 2
    assert (
        "投影片標題" in chunks[0].content
        and "欄位" in chunks[0].content
        and "數值" in chunks[0].content
    )


@pytest.mark.parametrize(
    "strategy", ["hierarchical", "fixed", "markdown-aware", "pdf-page", "cjk-sentence", "semantic"]
)
def test_pdf_citations_keep_physical_pages_for_every_strategy(strategy):
    chunks = get_chunker(strategy).chunk(
        "first page content\f\fthird page content",
        {"format": "pdf", "page_count": 3, "source_filename": "source.pdf"},
        {},
    )
    assert {chunk.metadata["page"] for chunk in chunks} == {1, 3}
    assert all(chunk.metadata["total_pages"] == 3 for chunk in chunks)
    assert all(chunk.metadata["source_filename"] == "source.pdf" for chunk in chunks)


@pytest.mark.parametrize("fault", ["unavailable", "initialization", "timeout", "corrupt", "empty"])
def test_docling_faults_use_independent_nonempty_native_parser(monkeypatch, fault):
    class Broken:
        def parse(self, path):
            if fault == "empty":
                return ParsedDocument(content="")
            if fault == "corrupt":
                raise ParseError.corrupt("解析失敗")
            if fault == "timeout":
                raise subprocess.TimeoutExpired("docling", 1)
            raise RemoteParseError.endpoint_unavailable("服務無法使用")

    def resolver(cls):
        if fault == "initialization":
            raise RuntimeError("initialization failed")
        return Broken()

    monkeypatch.setattr(ParserRegistry, "_get_docling_parser", classmethod(resolver))
    text, metadata, _ = extract_text("差旅.docx", _docx())
    assert "差旅規定" in text
    assert metadata["parser_fallback"] is True
    assert metadata["actual_parser"] == "DocxParser"
    assert metadata["quality_notice"]


def test_disabled_docling_docx_and_intact_table():
    text, metadata, _ = extract_text("差旅.docx", _docx())
    chunks = get_chunker("hierarchical").chunk(
        text, metadata, {"max_leaf_tokens": 2, "overlap_tokens": 0}
    )
    assert any("項目 | 金額\n交通費 | 依實報支" == c.content for c in chunks)
    assert metadata["actual_parser"] == "DocxParser"


def test_office_converter_missing_is_explicit(monkeypatch):
    monkeypatch.setattr(shutil, "which", lambda name: None)
    for name in ["old.doc", "old.ppt", "old.rtf"]:
        with pytest.raises(ParseError) as error:
            extract_text(name, b"{\\rtf1 plain text}")
        assert error.value.code == "E_PARSE_BAD_CONFIG"
        assert error.value.details["reason"] == "office_converter_missing"


def test_office_conversion_timeout_is_distinct(monkeypatch):
    monkeypatch.setattr(shutil, "which", lambda name: "/usr/bin/libreoffice")

    def timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired("libreoffice", 60)

    monkeypatch.setattr(subprocess, "run", timeout)
    with pytest.raises(ParseError) as error:
        extract_text("old.doc", b"old office")
    assert error.value.details["reason"] == "office_converter_timeout"


def test_converter_rejects_empty_output(monkeypatch):
    monkeypatch.setattr(shutil, "which", lambda name: "/usr/bin/libreoffice")
    monkeypatch.setattr(
        subprocess, "run", lambda *args, **kwargs: subprocess.CompletedProcess(args, 0)
    )
    with pytest.raises(ParseError) as error:
        extract_text("old.ppt", b"old office")
    assert error.value.details["reason"] == "office_conversion_failed"


def test_tesseract_missing_is_explicit_and_independent(monkeypatch):
    from anila_core.ingestion import local_ocr

    monkeypatch.setattr(shutil, "which", lambda name: None)
    with pytest.raises(ParseError) as error:
        local_ocr.image_text(b"pixels")
    assert error.value.details["reason"] == "offline_ocr_missing"
    assert "chi_tra" in error.value.user_message


def test_tesseract_missing_traditional_chinese_is_explicit(monkeypatch):
    from anila_core.ingestion import local_ocr

    monkeypatch.setattr(shutil, "which", lambda name: "/usr/bin/tesseract")
    monkeypatch.setattr(
        subprocess,
        "run",
        lambda *args, **kwargs: subprocess.CompletedProcess(args, 0, b"eng\n", b""),
    )
    with pytest.raises(ParseError) as error:
        local_ocr._binary()
    assert error.value.details["reason"] == "offline_ocr_languages_missing"


def test_failed_docling_image_uses_independent_local_ocr(monkeypatch):
    from anila_core.ingestion import image_ocr, local_ocr

    def broken(*args):
        raise TimeoutError("docling")

    monkeypatch.setattr(image_ocr, "_docling_image", broken)
    monkeypatch.setattr(local_ocr, "image_text", lambda data: "繁體中文辨識結果")
    metadata = {}
    assert image_ocr.recognize_image(b"pixels", object(), metadata) == "繁體中文辨識結果"
    assert metadata["actual_parser"] == "tesseract-local"


def test_scanned_pdf_uses_local_ocr_when_docling_disabled(monkeypatch, tmp_path):
    import fitz
    import pymupdf4llm
    from anila_core.ingestion import local_ocr

    path = tmp_path / "scan.pdf"
    with fitz.open() as pdf:
        pdf.new_page()
        pdf.save(path)
    monkeypatch.setattr(pymupdf4llm, "to_markdown", lambda *args, **kwargs: [{"text": ""}])
    monkeypatch.setattr(
        local_ocr,
        "pdf_document",
        lambda path: ParsedDocument(
            content="繁體中文掃描結果",
            format="pdf",
            metadata={"pages": 1, "parser": "tesseract-local", "ocr_used": True},
        ),
    )
    text, metadata, _ = extract_text("scan.pdf", path.read_bytes())
    assert text == "繁體中文掃描結果"
    assert metadata["actual_parser"] == "tesseract-local"


def test_empty_native_output_is_never_success(monkeypatch):
    class Empty:
        def parse(self, path):
            return ParsedDocument(content=" ")

    monkeypatch.setitem(ParserRegistry._PARSERS, ".txt", Empty())
    with pytest.raises(ParseError) as error:
        extract_text("blank.txt", b" ")
    assert error.value.code == "E_PARSE_CORRUPT"


def test_python_methods_keep_decorators_and_uploaded_source_is_never_executed():
    body = (
        'raise RuntimeError("source must remain data")\n'
        "class Policy:\n"
        "    @staticmethod\n"
        "    def lookup(value):\n"
        "        return value\n"
        "\n"
        "    def second(self):\n"
        "        return 2\n"
    )
    text, metadata, _ = extract_text("policy.py", body.encode())
    chunks = get_chunker("hierarchical").chunk(text, metadata, {"max_leaf_tokens": 1})
    assert "".join(chunk.content for chunk in chunks) == body
    methods = [chunk for chunk in chunks if chunk.metadata["symbol"].startswith("Policy.")]
    assert [chunk.metadata["symbol"] for chunk in methods] == ["Policy.lookup", "Policy.second"]
    assert "@staticmethod" in methods[0].content and "return value" in methods[0].content


def test_cpp_initializer_braces_keep_constructor_body_and_line_ranges():
    body = "class Policy {\nPolicy() : value{1}, other{2}\n{\n  ready();\n}\n};\n"
    text, metadata, _ = extract_text("policy.cpp", body.encode())
    chunks = get_chunker("fixed").chunk(text, metadata, {"size": 1, "overlap": 0})
    assert "".join(chunk.content for chunk in chunks) == body
    constructors = [chunk for chunk in chunks if chunk.metadata["symbol"] == "Policy"]
    assert len(constructors) == 1
    assert "value{1}, other{2}" in constructors[0].content
    assert "ready();" in constructors[0].content
    assert constructors[0].metadata["source_boundary_quality"] == "lexical"


def test_encrypted_pdf_has_specific_plain_language_error(tmp_path):
    import fitz

    source = tmp_path / "encrypted.pdf"
    with fitz.open() as pdf:
        pdf.new_page()
        pdf.save(
            source,
            encryption=fitz.PDF_ENCRYPT_AES_256,
            owner_pw="owner-fixture",
            user_pw="user-fixture",
        )
    with pytest.raises(ParseError) as error:
        extract_text("encrypted.pdf", source.read_bytes())
    assert error.value.details["reason"] == "encrypted_pdf"
    assert "密碼保護" in error.value.user_message


def test_independent_pdf_ocr_keeps_blank_page_slots_and_uses_traditional_chinese(
    monkeypatch, tmp_path
):
    import fitz
    from anila_core.ingestion import local_ocr

    source = tmp_path / "scan.pdf"
    with fitz.open() as pdf:
        for _ in range(3):
            pdf.new_page()
        pdf.save(source)

    monkeypatch.setattr(shutil, "which", lambda name: "/usr/bin/tesseract")
    calls = []

    def run(args, **kwargs):
        if args[1] == "--list-langs":
            return subprocess.CompletedProcess(args, 0, b"chi_tra\neng\n", b"")
        calls.append(args)
        content = "" if len(calls) == 2 else f"繁體中文第{len(calls)}頁"
        Path(args[2]).with_suffix(".txt").write_text(content, encoding="utf-8")
        return subprocess.CompletedProcess(args, 0, b"", b"")

    monkeypatch.setattr(subprocess, "run", run)
    parsed = local_ocr.pdf_document(str(source))
    assert parsed.content.split("\f") == ["繁體中文第1頁", "", "繁體中文第3頁"]
    assert parsed.metadata["pages"] == 3
    assert parsed.metadata["actual_parser"] == "tesseract-local"
    assert parsed.metadata["quality"] == "ocr_text_only"
    assert all(args[3:5] == ["-l", "chi_tra+eng"] for args in calls)


def test_standalone_ocr_preserves_image_for_existing_image_retrieval(monkeypatch):
    from anila_core.ingestion import image_ocr

    def recognize(data, parser, metadata):
        metadata["actual_parser"] = "tesseract-local"
        return "繁體中文辨識結果"

    monkeypatch.setattr(image_ocr, "recognize_image", recognize)
    text, metadata, images = extract_text("scan.png", b"test-pixels")
    assert "繁體中文辨識結果" in text
    assert len(images) == 1 and next(iter(images.values())).image_bytes == b"test-pixels"
    assert metadata["actual_parser"] == "tesseract-local"


def test_standalone_empty_ocr_never_becomes_placeholder_only_success(monkeypatch):
    from anila_core.ingestion import image_ocr

    def recognize(data, parser, metadata):
        metadata["actual_parser"] = "tesseract-local"
        return " "

    monkeypatch.setattr(image_ocr, "recognize_image", recognize)
    with pytest.raises(ParseError) as error:
        extract_text("scan.png", b"test-pixels")
    assert error.value.details["reason"] == "empty_image_ocr_output"


def _png() -> bytes:
    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (1, 1), "white").save(buffer, "PNG")
    return buffer.getvalue()


@pytest.mark.parametrize(
    "markers",
    [
        "[[IMAGE:remote_scan]]",
        "\f<!-- page_break -->\f[[PAGE:2]]",
        "![scanned page](scan.png)",
        '<img src="scan.png">',
        "# Page 1\n[[IMAGE:scan]]\n | |\n|---|---|",
    ],
)
def test_marker_only_docling_image_uses_independent_ocr(monkeypatch, markers):
    from anila_core.ingestion import local_ocr

    class Remote:
        def parse(self, path):
            return ParsedDocument(content=markers, format="pdf")

    monkeypatch.setattr(ParserRegistry, "_get_docling_parser", classmethod(lambda cls: Remote()))
    called = []

    def recognize(data):
        called.append(data)
        return "繁體中文實際辨識結果"

    monkeypatch.setattr(local_ocr, "image_text", recognize)
    pixels = _png()
    text, metadata, _ = extract_text("scan.png", pixels)
    assert "繁體中文實際辨識結果" in text
    assert called == [pixels]
    assert metadata["actual_parser"] == "tesseract-local"


def test_marker_only_remote_pdf_with_images_does_not_bypass_local_ocr(monkeypatch, tmp_path):
    import base64
    import fitz
    import pymupdf4llm
    from anila_core.ingestion import local_ocr
    from anila_core.ingestion.docling_parser import RemoteDoclingParser

    source = tmp_path / "scan.pdf"
    with fitz.open() as pdf:
        pdf.new_page().insert_image(fitz.Rect(0, 0, 100, 100), stream=_png())
        pdf.save(source)

    class Remote:
        def parse(self, path):
            return RemoteDoclingParser._reconstruct(
                self,
                {
                    "markdown": "",
                    "page_count": 1,
                    "images": [
                        {
                            "id": "remote_scan",
                            "png_b64": base64.b64encode(_png()).decode("ascii"),
                            "page": 1,
                        }
                    ],
                },
                path,
                ".pdf",
            )

    monkeypatch.setattr(ParserRegistry, "_get_docling_parser", classmethod(lambda cls: Remote()))
    monkeypatch.setattr(pymupdf4llm, "to_markdown", lambda *args, **kwargs: [{"text": ""}])
    calls = []

    def recognize(path):
        calls.append(path)
        return ParsedDocument(
            content="繁體中文掃描頁",
            format="pdf",
            metadata={"pages": 1, "parser": "tesseract-local"},
        )

    monkeypatch.setattr(local_ocr, "pdf_document", recognize)
    text, metadata, _ = extract_text("scan.pdf", source.read_bytes())
    assert text == "繁體中文掃描頁" and len(calls) == 1
    assert metadata["actual_parser"] == "tesseract-local"
    assert metadata["parser_fallback"] is True


def test_marker_only_both_ocr_engines_fail_explicitly(monkeypatch):
    from anila_core.ingestion import local_ocr

    class Remote:
        def parse(self, path):
            return ParsedDocument(content="[[IMAGE:scan]]", format="pdf")

    monkeypatch.setattr(ParserRegistry, "_get_docling_parser", classmethod(lambda cls: Remote()))
    monkeypatch.setattr(local_ocr, "image_text", lambda data: "Page 1\n[[IMAGE:local_scan]]")
    with pytest.raises(ParseError) as error:
        extract_text("scan.png", _png())
    assert error.value.details["reason"] == "empty_image_ocr_output"


def test_docx_table_cell_image_is_recognized_and_table_stays_intact(monkeypatch):
    from docx import Document
    from anila_core.ingestion import local_ocr

    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).paragraphs[0].add_run().add_picture(io.BytesIO(_png()))
    content = io.BytesIO()
    document.save(content)
    calls = []

    def recognize(data):
        calls.append(data)
        return "表格內的繁體中文"

    monkeypatch.setattr(local_ocr, "image_text", recognize)
    text, metadata, images = extract_text("table-scan.docx", content.getvalue())
    assert "表格內的繁體中文" in text and len(calls) == len(images) == 1
    chunks = get_chunker("hierarchical").chunk(
        text, metadata, {"max_leaf_tokens": 1, "overlap_tokens": 0}
    )
    assert any(chunk.content == metadata["table_blocks"][0] for chunk in chunks)


def test_empty_docx_table_is_not_recovered_text():
    from docx import Document

    document = Document()
    document.add_table(rows=1, cols=2)
    content = io.BytesIO()
    document.save(content)
    with pytest.raises(ParseError) as error:
        extract_text("empty-table.docx", content.getvalue())
    assert error.value.details["reason"] == "empty_parser_output"


def test_unclosed_fence_scans_each_line_once(monkeypatch):
    from anila_core.ingestion.chunking_plugins import structure

    original = structure._fence
    calls = 0

    def counted(line):
        nonlocal calls
        calls += 1
        return original(line)

    monkeypatch.setattr(structure, "_fence", counted)
    body = "```python\nx\n" * 4000
    blocks = structure._structural_blocks(body)
    assert len(blocks) == 1 and blocks[0].group() == body
    assert calls == len(body.splitlines())
    from anila_core.ingestion.errors import ChunkError
    with pytest.raises(ChunkError) as error:
        get_chunker("fixed").chunk(body, {}, {"size": 1, "overlap": 0})
    assert error.value.code == "E_CHUNK_TOO_LARGE"
    masked = structure.mask_fenced_code(body)
    assert masked == "".join(char if char in "\r\n" else " " for char in body)
    with pytest.raises(ChunkError):
        get_chunker("markdown-aware").chunk(body, {}, {"max_leaf_tokens": 100_000})


@pytest.mark.skipif(
    not (shutil.which("libreoffice") or shutil.which("soffice")),
    reason="existing LibreOffice binary unavailable",
)
def test_real_offline_rtf_conversion():
    text, metadata, _ = extract_text(
        "legacy.rtf", b"{\\rtf1\\ansi Offline travel policy\\par Article 12: reimbursement.}"
    )
    assert "Offline travel policy" in text and "Article 12" in text
    assert metadata["conversion"] == "libreoffice-local"
    assert metadata["format"] == "rtf"


def _export_legacy(source: Path, extension: str, work: Path) -> Path:
    binary = shutil.which("libreoffice") or shutil.which("soffice")
    output = work / "legacy"
    output.mkdir()
    result = subprocess.run(
        [
            binary,
            "-env:UserInstallation=" + (work / "export-profile").as_uri(),
            "--headless",
            "--nologo",
            "--nodefault",
            "--nofirststartwizard",
            "--convert-to",
            extension,
            "--outdir",
            str(output),
            str(source),
        ],
        capture_output=True,
        timeout=60,
        check=False,
    )
    target = output / (source.stem + "." + extension)
    assert result.returncode == 0 and target.is_file(), "generated legacy fixture conversion failed"
    return target


@pytest.mark.skipif(
    not (shutil.which("libreoffice") or shutil.which("soffice")),
    reason="existing LibreOffice binary unavailable",
)
def test_real_offline_doc_roundtrip(tmp_path):
    source = tmp_path / "travel.docx"
    source.write_bytes(_docx())
    legacy = _export_legacy(source, "doc", tmp_path)
    text, metadata, _ = extract_text("差旅.doc", legacy.read_bytes())
    assert "差旅規定" in text and "交通費" in text and "依實報支" in text
    assert metadata["conversion"] == "libreoffice-local"
    assert metadata["format"] == "doc"


@pytest.mark.skipif(
    not (shutil.which("libreoffice") or shutil.which("soffice")),
    reason="existing LibreOffice binary unavailable",
)
def test_real_offline_ppt_roundtrip(tmp_path):
    source = tmp_path / "travel.odp"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("mimetype", "application/vnd.oasis.opendocument.presentation")
        archive.writestr(
            "META-INF/manifest.xml",
            '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.presentation"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>',
        )
        archive.writestr(
            "content.xml",
            '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2"><office:body><office:presentation><draw:page draw:name="page1"><draw:frame><draw:text-box><text:p>Offline travel slides</text:p></draw:text-box></draw:frame></draw:page></office:presentation></office:body></office:document-content>',
        )
    legacy = _export_legacy(source, "ppt", tmp_path)
    text, metadata, _ = extract_text("travel.ppt", legacy.read_bytes())
    assert "Offline travel slides" in text
    assert metadata["conversion"] == "libreoffice-local"
    assert metadata["slide_count"] == 1
    assert metadata["format"] == "ppt"
