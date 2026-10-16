"""OOXML markers are evidence of declared/cached breaks, not physical pages."""
import pytest

from anila_core.ingestion.parser_registry import DocxParser
from anila_core.ingestion.chunking_plugins import get_chunker

docx = pytest.importorskip('docx')
from docx.oxml import OxmlElement  # noqa: E402
from docx.enum.text import WD_BREAK  # noqa: E402


def parse(document, tmp_path):
    path = tmp_path / 'synthetic.docx'
    document.save(path)
    return DocxParser().parse(str(path))


def test_explicit_and_cached_markers_keep_sources_without_physical_page_numbers(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph('第一段')
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    paragraph.add_run('第二段')
    rendered = OxmlElement('w:lastRenderedPageBreak')
    paragraph.add_run()._r.append(rendered)
    paragraph.add_run('第三段')
    parsed = parse(document, tmp_path)
    assert parsed.metadata['docx_page_marker_counts'] == {'explicit_page_break': 1, 'last_rendered_page_break': 1}
    events = parsed.metadata['docx_page_boundary_evidence']
    assert [(event['source'], event['confidence']) for event in events] == [
        ('w:br[type=page]', 'declared_break_not_physical_page'),
        ('w:lastRenderedPageBreak', 'cached_layout_not_verified'),
    ]
    assert [event['source_text_char_offset'] for event in events] == [3, 6]
    assert parsed.metadata['page_location_quality'] == 'mixed_markers_no_physical_pages'
    assert not {'pages', 'page_count', 'page', 'total_pages'} & parsed.metadata.keys()
    assert all(text in parsed.content for text in ('第一段', '第二段', '第三段'))


def test_cached_marker_alone_is_not_verified_pagination(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph('快取前')
    paragraph.add_run()._r.append(OxmlElement('w:lastRenderedPageBreak'))
    paragraph.add_run('快取後')
    parsed = parse(document, tmp_path)
    assert parsed.metadata['page_location_quality'] == 'cached_markers_no_physical_pages'
    assert parsed.metadata['docx_page_marker_counts'] == {'last_rendered_page_break': 1}
    assert 'page_count' not in parsed.metadata


def test_no_marker_and_inherited_style_do_not_estimate_pages(tmp_path):
    document = docx.Document()
    document.styles['Normal'].paragraph_format.page_break_before = True
    document.add_paragraph('樣式繼承須排版確認')
    parsed = parse(document, tmp_path)
    assert parsed.metadata['page_location_quality'] == 'unknown'
    assert parsed.metadata['docx_page_marker_counts'] == {}
    assert parsed.metadata['docx_page_boundary_evidence'] == []


def test_explicit_section_type_retains_source_not_odd_even_page_estimate(tmp_path):
    document = docx.Document()
    paragraph = document.add_paragraph('節界前')
    section = OxmlElement('w:sectPr')
    section_type = OxmlElement('w:type')
    section_type.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', 'oddPage')
    section.append(section_type)
    paragraph._p.get_or_add_pPr().append(section)
    parsed = parse(document, tmp_path)
    assert parsed.metadata['docx_page_marker_counts'] == {'section_page_break': 1}
    assert 'oddPage' in parsed.metadata['docx_page_boundary_evidence'][0]['source']
    assert not {'pages', 'page_count'} & parsed.metadata.keys()


def test_page_break_before_is_declared_and_false_values_are_ignored(tmp_path):
    document = docx.Document()
    document.add_paragraph('前段')
    document.add_paragraph('宣告分頁').paragraph_format.page_break_before = True
    document.add_paragraph('不分頁').paragraph_format.page_break_before = False
    parsed = parse(document, tmp_path)
    assert parsed.metadata['docx_page_marker_counts'] == {'page_break_before': 1}
    assert parsed.metadata['docx_page_boundary_evidence'][0]['location'] == {'body_block': 1}
    assert parsed.metadata['page_location_quality'] == 'declared_breaks_no_physical_pages'


def test_merged_table_marker_is_not_double_counted(tmp_path):
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    cell = table.cell(0, 0).merge(table.cell(0, 1))
    cell.paragraphs[0].add_run('表格前').add_break(WD_BREAK.PAGE)
    cell.paragraphs[0].add_run('表格後')
    parsed = parse(document, tmp_path)
    assert parsed.metadata['docx_page_marker_counts'] == {'explicit_page_break': 1}
    assert parsed.metadata['docx_page_boundary_evidence'][0]['location'] == {
        'body_block': 0, 'table_row': 0, 'table_cell': 0, 'paragraph': 0,
    }


def test_marker_evidence_is_bounded_but_counts_are_complete(tmp_path):
    document = docx.Document()
    for index in range(140):
        document.add_paragraph(str(index)).add_run().add_break(WD_BREAK.PAGE)
    parsed = parse(document, tmp_path)
    assert parsed.metadata['docx_page_marker_counts'] == {'explicit_page_break': 140}
    assert len(parsed.metadata['docx_page_boundary_evidence']) == 128
    assert parsed.metadata['docx_page_boundary_evidence_truncated'] is True


def test_pdf_page_strategy_cannot_create_physical_pages_for_unrendered_docx():
    chunks = get_chunker('pdf-page').chunk('未排版的 Word 文字', {'format': 'docx'}, {})
    assert chunks
    assert all(not {'page', 'total_pages'} & chunk.metadata.keys() for chunk in chunks)


def test_shared_extraction_keeps_marker_provenance_without_page_count(tmp_path, monkeypatch):
    from anila_core.ingestion.parser_registry import ParserRegistry
    from anila_core.ingestion.parsers import extract_text

    monkeypatch.setattr(ParserRegistry, '_get_docling_parser', classmethod(lambda cls: None))
    document = docx.Document()
    document.add_paragraph('共享入口前').add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph('共享入口後')
    path = tmp_path / 'shared.docx'
    document.save(path)
    text, metadata, _ = extract_text('shared.docx', path.read_bytes())
    assert '共享入口前' in text and '共享入口後' in text
    assert metadata['source_filename'] == 'shared.docx'
    assert metadata['docx_page_marker_counts'] == {'explicit_page_break': 1}
    assert metadata['has_page_boundaries'] is False
    assert 'page_count' not in metadata


@pytest.mark.parametrize('strategy', ['fixed', 'markdown-aware', 'hierarchical', 'pdf-page', 'cjk-sentence', 'semantic'])
def test_docx_location_notice_reaches_chunks_without_invented_pages(tmp_path, strategy):
    document = docx.Document()
    document.add_paragraph('第一段正文')
    document.add_paragraph('第二段正文').paragraph_format.page_break_before = True
    parsed = parse(document, tmp_path)
    params = {'_segments': [parsed.content], '_embeddings': [[1.0, 0.0]]} if strategy == 'semantic' else {}
    chunks = get_chunker(strategy).chunk(parsed.content, {**parsed.metadata, 'format': 'docx'}, params)
    leaves = [chunk for chunk in chunks if chunk.metadata.get('chunk_type') != 'parent']
    assert leaves
    for chunk in leaves:
        assert chunk.metadata['page_location_source'] == 'docx_ooxml_markers'
        assert chunk.metadata['page_location_quality'] == 'declared_breaks_no_physical_pages'
        assert chunk.metadata['page_location_notice']
        assert not {'page', 'total_pages'} & chunk.metadata.keys()
        assert 'docx_page_boundary_evidence' not in chunk.metadata
