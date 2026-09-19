"""DOCX tables must land at their true document position, not all at the end.

Pins: paragraphs and tables interleave in body order; a table's lead-in
sentence stays adjacent to its table; ``[[IMAGE:<id>]]`` placeholders stay
with the paragraph they came from; and the reordering loses nothing — the
set of emitted blocks is identical to what the old two-collection walk
produced.

Why two tables in every order assertion: with a single table, "all tables
at the end" and "correct position" can coincide, and "reversed" is
invisible.  Two tables kill both mutations.

Also pins the *absence* face: ``w:sdt`` (content control) and ``w:ins``
(tracked insertion) wrappers hide their paragraphs and tables from this
parser and always have.  That behaviour is unchanged, but it is now
counted and logged rather than silent, so both the drop and the warning
are pinned here — the day someone teaches the walk to descend, a test
moves instead of the output quietly changing.
"""

from __future__ import annotations

import io
import logging
import re

import pytest

import anila_core.ingestion.parser_registry as parser_registry
from anila_core.ingestion.parser_registry import DocxParser

docx = pytest.importorskip("docx", reason="python-docx is part of the [rag] extra")

# Header/footer ``add_table`` demands an explicit width (unlike the body's).
from docx.shared import Inches  # noqa: E402  — must follow the importorskip


# ──────────────────────────────────────────────────────────────────────
# helpers
# ──────────────────────────────────────────────────────────────────────

def _save(document, tmp_path, name: str = "sample.docx") -> str:
    path = tmp_path / name
    document.save(str(path))
    return str(path)


def _blocks(content: str) -> list[str]:
    """Split parsed content back into the blocks the parser joined."""
    return content.split("\n\n")


def _fill(table, rows: list[list[str]]) -> None:
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value


def _tiny_png() -> io.BytesIO:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (4, 4), "red").save(buf, "PNG")
    buf.seek(0)
    return buf


def _old_order_blocks(self, doc, skipped_with_content=None):
    """The pre-fix walk: every paragraph, then every table.

    Kept in the tests (not in production) so the content-equality test can
    compare against the exact behaviour the fix replaced, using the same
    rendering code for both. Extra arg is unused: production now threads
    the skipped-wrapper counter through this method.
    """
    for para in doc.paragraphs:
        yield "paragraph", para
    for table in doc.tables:
        yield "table", table


def _normalise_image_ids(blocks: list[str]) -> list[str]:
    """Image ids are fresh uuids per parse; compare shape, not identity."""
    return [re.sub(r"\[\[IMAGE:[^\]]+\]\]", "[[IMAGE:*]]", b) for b in blocks]


_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# python-docx has no API for content controls or tracked changes, so these
# body children have to be injected as raw XML — which is exactly how Word
# writes them when a document is built from a content-control template.
_SDT_PARAGRAPH = (
    "<w:sdt><w:sdtContent>"
    "<w:p><w:r><w:t>SDT PARA</w:t></w:r></w:p>"
    "</w:sdtContent></w:sdt>"
)
_SDT_TABLE = (
    "<w:sdt><w:sdtContent><w:tbl><w:tr><w:tc>"
    "<w:p><w:r><w:t>SDT CELL</w:t></w:r></w:p>"
    "</w:tc></w:tr></w:tbl></w:sdtContent></w:sdt>"
)
_INS_PARAGRAPH = (
    '<w:ins w:id="9" w:author="tester" w:date="2026-08-06T00:00:00Z">'
    "<w:p><w:r><w:t>INS PARA</w:t></w:r></w:p>"
    "</w:ins>"
)


def _inject(document, xml: str) -> None:
    """Append a raw body-level element, keeping the trailing ``w:sectPr`` last."""
    from lxml import etree

    element = etree.fromstring(f'<w:root xmlns:w="{_W}">{xml}</w:root>')[0]
    body = document.element.body
    sect_pr = body.find(f"{{{_W}}}sectPr")
    if sect_pr is not None:
        sect_pr.addprevious(element)
    else:
        body.append(element)


def _wrapper_warnings(caplog) -> list[str]:
    return [
        record.getMessage()
        for record in caplog.records
        if record.levelno == logging.WARNING
        and record.name == parser_registry.__name__
        and "not top-level w:p / w:tbl" in record.getMessage()
    ]


# ──────────────────────────────────────────────────────────────────────
# order
# ──────────────────────────────────────────────────────────────────────

def test_tables_appear_at_their_document_position(tmp_path) -> None:
    document = docx.Document()
    document.add_paragraph("para A")
    _fill(document.add_table(rows=1, cols=2), [["a1", "b1"]])
    document.add_paragraph("para B")
    _fill(document.add_table(rows=1, cols=2), [["a2", "b2"]])
    document.add_paragraph("para C")

    parsed = DocxParser().parse(_save(document, tmp_path))

    assert _blocks(parsed.content) == [
        "para A",
        "a1 | b1",
        "para B",
        "a2 | b2",
        "para C",
    ]


def test_lead_in_sentence_stays_adjacent_to_its_table(tmp_path) -> None:
    """The failure the users actually hit: sentence and table in one chunk."""
    document = docx.Document()
    document.add_paragraph("推進段各項規格如下表：")
    _fill(document.add_table(rows=1, cols=2), [["推力", "50 kN"]])
    document.add_paragraph("導引段各項規格如下表：")
    _fill(document.add_table(rows=1, cols=2), [["精度", "3 m"]])

    blocks = _blocks(DocxParser().parse(_save(document, tmp_path)).content)

    assert blocks.index("推力 | 50 kN") == blocks.index("推進段各項規格如下表：") + 1
    assert blocks.index("精度 | 3 m") == blocks.index("導引段各項規格如下表：") + 1


def test_multi_row_table_keeps_its_rendering_and_its_position(tmp_path) -> None:
    document = docx.Document()
    document.add_paragraph("intro")
    _fill(
        document.add_table(rows=2, cols=2),
        [["h1", "h2"], ["v1", "v2"]],
    )
    document.add_paragraph("middle")
    _fill(document.add_table(rows=1, cols=3), [["x", "y", "z"]])
    document.add_paragraph("outro")

    assert _blocks(DocxParser().parse(_save(document, tmp_path)).content) == [
        "intro",
        "h1 | h2\nv1 | v2",
        "middle",
        "x | y | z",
        "outro",
    ]


def test_headings_interleave_with_tables_and_title_is_first_heading(tmp_path) -> None:
    document = docx.Document()
    _fill(document.add_table(rows=1, cols=1), [["前言表"]])
    document.add_paragraph("系統規格書", style="Heading 1")
    _fill(document.add_table(rows=1, cols=1), [["表一"]])
    document.add_paragraph("推進段", style="Heading 2")
    _fill(document.add_table(rows=1, cols=1), [["表二"]])

    parsed = DocxParser().parse(_save(document, tmp_path))

    assert _blocks(parsed.content) == [
        "前言表",
        "# 系統規格書",
        "表一",
        "## 推進段",
        "表二",
    ]
    assert parsed.metadata["title"] == "系統規格書"


def test_table_before_any_paragraph_is_emitted_first(tmp_path) -> None:
    document = docx.Document()
    _fill(document.add_table(rows=1, cols=1), [["leading table"]])
    document.add_paragraph("after")
    _fill(document.add_table(rows=1, cols=1), [["trailing table"]])

    assert _blocks(DocxParser().parse(_save(document, tmp_path)).content) == [
        "leading table",
        "after",
        "trailing table",
    ]


# ──────────────────────────────────────────────────────────────────────
# nothing is lost
# ──────────────────────────────────────────────────────────────────────

def _edge_document():
    document = docx.Document()
    document.add_paragraph("標題", style="Heading 1")
    document.add_paragraph("")                       # empty paragraph, dropped
    document.add_paragraph("plain")
    _fill(document.add_table(rows=1, cols=2), [["", ""]])   # blank cells
    document.add_table(rows=0, cols=2)                      # no rows at all
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "outer cell"
    _fill(outer.cell(0, 0).add_table(rows=1, cols=1), [["nested"]])
    document.add_paragraph("tail")
    document.sections[0].header.add_table(rows=1, cols=1, width=Inches(1))
    return document


def test_reordering_loses_no_block_versus_old_collection_walk(
    tmp_path, monkeypatch
) -> None:
    """Same set of blocks as the pre-fix walk — only the order changed."""
    path = _save(_edge_document(), tmp_path)

    new_blocks = _normalise_image_ids(_blocks(DocxParser().parse(path).content))

    monkeypatch.setattr(DocxParser, "_iter_body_blocks", _old_order_blocks)
    old_blocks = _normalise_image_ids(_blocks(DocxParser().parse(path).content))

    assert sorted(new_blocks) == sorted(old_blocks)
    assert new_blocks != old_blocks, "fixture must actually exercise reordering"


def test_empty_table_is_skipped_and_blank_cell_table_is_kept(tmp_path) -> None:
    document = docx.Document()
    document.add_paragraph("head")
    _fill(document.add_table(rows=1, cols=2), [["", ""]])
    document.add_table(rows=0, cols=2)
    document.add_paragraph("tail")

    assert _blocks(DocxParser().parse(_save(document, tmp_path)).content) == [
        "head",
        " | ",
        "tail",
    ]


def test_nested_table_text_is_not_emitted(tmp_path) -> None:
    """Documented gap, unchanged by this fix.

    ``doc.tables`` never returned nested tables and ``_Cell.text`` only joins
    the cell's own paragraphs, so nested table text was dropped before the
    fix too.  Pinned here so a later change to the walk cannot silently start
    or stop emitting it without this test speaking up.
    """
    document = docx.Document()
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "outer cell"
    _fill(outer.cell(0, 0).add_table(rows=1, cols=1), [["nested value"]])

    content = DocxParser().parse(_save(document, tmp_path)).content

    assert "outer cell" in content
    assert "nested value" not in content


def test_header_and_footer_tables_are_not_emitted(tmp_path) -> None:
    """Documented gap, unchanged: the walk covers the body only."""
    document = docx.Document()
    document.add_paragraph("body para")
    _fill(document.sections[0].header.add_table(rows=1, cols=1, width=Inches(1)),
          [["header table"]])
    _fill(document.sections[0].footer.add_table(rows=1, cols=1, width=Inches(1)),
          [["footer table"]])

    content = DocxParser().parse(_save(document, tmp_path)).content

    assert "body para" in content
    assert "header table" not in content
    assert "footer table" not in content


# ──────────────────────────────────────────────────────────────────────
# image placeholders
# ──────────────────────────────────────────────────────────────────────

def test_image_placeholder_stays_with_its_paragraph_across_tables(tmp_path) -> None:
    document = docx.Document()
    document.add_paragraph("before")
    _fill(document.add_table(rows=1, cols=1), [["table one"]])
    para = document.add_paragraph("圖示如下")
    para.add_run().add_picture(_tiny_png())
    _fill(document.add_table(rows=1, cols=1), [["table two"]])
    document.add_paragraph("after")

    parsed = DocxParser().parse(_save(document, tmp_path))

    assert len(parsed.images) == 1
    image_id = next(iter(parsed.images))
    assert _blocks(parsed.content) == [
        "before",
        "table one",
        "圖示如下",
        f"[[IMAGE:{image_id}]]",
        "table two",
        "after",
    ]


# ──────────────────────────────────────────────────────────────────────
# the silent face: content controls and tracked changes
#
# These pin the CURRENT behaviour, which is also the pre-fix behaviour:
# a w:sdt / w:ins wrapper hides its paragraphs and tables from the walk.
# Nothing here is a regression — it is a long-standing gap that had no
# test watching it, on exactly the template shape an organisation
# standardises on.  Pinned so that changing it has to be deliberate.
# ──────────────────────────────────────────────────────────────────────

def test_sdt_wrapped_paragraph_is_dropped(tmp_path) -> None:
    document = docx.Document()
    document.add_paragraph("visible para")
    _inject(document, _SDT_PARAGRAPH)
    document.add_paragraph("trailing para")

    content = DocxParser().parse(_save(document, tmp_path)).content

    assert "visible para" in content
    assert "trailing para" in content
    assert "SDT PARA" not in content


def test_sdt_wrapped_table_is_dropped(tmp_path) -> None:
    document = docx.Document()
    document.add_paragraph("visible para")
    _fill(document.add_table(rows=1, cols=1), [["REAL CELL"]])
    _inject(document, _SDT_TABLE)

    content = DocxParser().parse(_save(document, tmp_path)).content

    assert "REAL CELL" in content
    assert "SDT CELL" not in content


def test_ins_wrapped_paragraph_is_dropped(tmp_path) -> None:
    document = docx.Document()
    document.add_paragraph("visible para")
    _inject(document, _INS_PARAGRAPH)

    content = DocxParser().parse(_save(document, tmp_path)).content

    assert "visible para" in content
    assert "INS PARA" not in content


def test_dropped_wrappers_do_not_disturb_the_order_of_what_survives(tmp_path) -> None:
    """A content-control template still gets its real tables in position."""
    document = docx.Document()
    document.add_paragraph("para A")
    _inject(document, _SDT_PARAGRAPH)
    _fill(document.add_table(rows=1, cols=2), [["a1", "b1"]])
    _inject(document, _SDT_TABLE)
    document.add_paragraph("para B")
    _fill(document.add_table(rows=1, cols=2), [["a2", "b2"]])

    assert _blocks(DocxParser().parse(_save(document, tmp_path)).content) == [
        "para A",
        "a1 | b1",
        "para B",
        "a2 | b2",
    ]


# ──────────────────────────────────────────────────────────────────────
# the silent face gets a voice
# ──────────────────────────────────────────────────────────────────────

def test_skipped_content_wrappers_are_counted_in_one_warning(
    tmp_path, caplog
) -> None:
    document = docx.Document()
    document.add_paragraph("visible para")
    _inject(document, _SDT_PARAGRAPH)
    _inject(document, _SDT_TABLE)
    _inject(document, _INS_PARAGRAPH)

    with caplog.at_level(logging.WARNING, logger=parser_registry.__name__):
        parsed = DocxParser().parse(_save(document, tmp_path))

    warnings = _wrapper_warnings(caplog)
    assert len(warnings) == 1, "one aggregate warning per document, not one per element"
    assert "skipped 3 body element(s)" in warnings[0]
    assert "w:sdt=2" in warnings[0]
    assert "w:ins=1" in warnings[0]
    assert parsed.metadata["docx_wrapped_skipped"] == 3


def test_plain_document_logs_no_wrapper_warning(tmp_path, caplog) -> None:
    """The must-not-cry control: a normal document stays silent."""
    document = docx.Document()
    document.add_paragraph("標題", style="Heading 1")
    document.add_paragraph("para A")
    _fill(document.add_table(rows=2, cols=2), [["h1", "h2"], ["v1", "v2"]])
    document.add_paragraph("para B")
    _fill(document.add_table(rows=1, cols=1), [[""]])
    document.add_table(rows=0, cols=2)

    with caplog.at_level(logging.WARNING, logger=parser_registry.__name__):
        parsed = DocxParser().parse(_save(document, tmp_path))

    assert _wrapper_warnings(caplog) == []
    assert "docx_wrapped_skipped" not in parsed.metadata


def test_inert_body_children_log_no_wrapper_warning(tmp_path, caplog) -> None:
    """w:sectPr, bookmarks, comments and PIs carry no text — they must not cry.

    Detection is by content (does this child contain a w:p / w:tbl?), not by
    a list of known-inert tags, so this is the closing half of that rule.
    """
    document = docx.Document()
    document.add_paragraph("alpha")
    _inject(document, '<w:bookmarkStart w:id="1" w:name="mark"/>')
    _inject(document, '<w:bookmarkEnd w:id="1"/>')
    _inject(document, "<w:permStart/>")
    _inject(document, "<!-- an xml comment -->")
    _inject(document, "<?some-processing-instruction value?>")
    document.add_paragraph("omega")

    with caplog.at_level(logging.WARNING, logger=parser_registry.__name__):
        parsed = DocxParser().parse(_save(document, tmp_path))

    assert _blocks(parsed.content) == ["alpha", "omega"]
    assert _wrapper_warnings(caplog) == []
