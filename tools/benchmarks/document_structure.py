"""Generate synthetic local documents and measure parser/chunker resource use.

Run from /tmp with env -i and python -I -S. No network, subprocess or model calls.
"""
# ruff: noqa: E402 -- the audit guard must precede dependency imports.

import os
import sys


def _guard(event, args):
    if event == "open" and args and isinstance(args[0], (str, bytes, os.PathLike)):
        path = os.fsdecode(args[0])
        parts = path.replace("\\", "/").split("/")
        if "secrets" in parts or parts[-1].startswith(".env") or parts[-1].endswith((".pem", ".key")):
            raise PermissionError("protected_file")
        flags = args[2]
        writing = isinstance(flags, int) and flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC | os.O_APPEND)
        if writing and not (os.path.abspath(path).startswith("/tmp/") or path == "/dev/null"):
            raise PermissionError("write_outside_tmp")
    if event in {"socket.connect", "socket.getaddrinfo", "socket.bind", "socket.sendto", "subprocess.Popen", "os.system", "os.posix_spawn", "os.fork"}:
        raise PermissionError("external_io_disabled")


sys.addaudithook(_guard)
if os.getcwd() != "/tmp":
    raise SystemExit("Run from /tmp")

import argparse
import hashlib
import json
from pathlib import Path
import resource
import sysconfig
import time
import tracemalloc

sys.path.extend([sysconfig.get_path("purelib"), sysconfig.get_path("purelib", scheme="posix_user")])
REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO / "packages/anila-core/src"))

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
import fitz

from anila_core.ingestion.chunking_plugins import get_chunker
from anila_core.ingestion.local_ocr import image_text
from anila_core.ingestion.parser_registry import DocxParser
from anila_core.ingestion.source_code import source_units, summarize_boundary_quality


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    output = Path(args.output).resolve()
    if not str(output).startswith("/tmp/"):
        raise SystemExit("Output must be under /tmp")
    output.mkdir(parents=True, exist_ok=False)
    document = Document()
    document.add_heading("合成文件頁界", level=1)
    paragraph = document.add_paragraph("明確頁界前")
    paragraph.add_run().add_break(WD_BREAK.PAGE)
    paragraph.add_run("明確頁界後")
    paragraph.add_run()._r.append(OxmlElement("w:lastRenderedPageBreak"))
    paragraph.add_run("快取頁界後")
    table = document.add_table(rows=3, cols=2)
    for index, row in enumerate(table.rows):
        row.cells[0].text, row.cells[1].text = f"項目{index}", f"內容{index}"
    document.save(output / "markers.docx")
    automatic = Document()
    for index in range(60):
        automatic.add_paragraph(f"Paragraph {index}. " + "Synthetic pagination evidence. " * 20)
    automatic.save(output / "automatic.docx")
    with fitz.open() as pdf:
        for index in range(3):
            page = pdf.new_page()
            page.insert_text((72, 72), f"Synthetic physical PDF page {index + 1}")
        pdf.save(output / "physical-three-pages.pdf")
    with fitz.open() as pdf:
        page = pdf.new_page(width=600, height=200)
        page.insert_text((20, 55), "繁體中文掃描合成樣本", fontname="china-t", fontsize=24)
        page.insert_text((20, 110), "ANILA synthetic OCR 2026", fontsize=20)
        page.get_pixmap(dpi=150, alpha=False).save(output / "scan-chi-tra.png")
    (output / "scan-chi-tra.expected.txt").write_text("繁體中文掃描合成樣本\nANILA synthetic OCR 2026\n")
    source = "".join(f"int function_{index}() {{ return 1'000; }}\n" for index in range(400))
    (output / "functions.cpp").write_text(source)
    large_source = "".join(f"int function_{index}() {{ return 1'000; }}\n" for index in range(2000))
    (output / "functions-large.cpp").write_text(large_source)

    def docx_case(path):
        parsed = DocxParser().parse(str(path))
        metadata = {**parsed.metadata, "format": "docx"}
        chunks = get_chunker("pdf-page").chunk(parsed.content, metadata, {})
        return {"text_chars": len(parsed.content), "chunks": len(chunks), "page_location_quality": metadata["page_location_quality"], "marker_counts": metadata["docx_page_marker_counts"], "physical_page_fields": any("page" in chunk.metadata or "total_pages" in chunk.metadata for chunk in chunks)}

    def cpp_case(text, functions):
        units = source_units(text, "cpp")
        chunks = get_chunker("fixed").chunk(text, {"source_units": units, **summarize_boundary_quality(units)}, {})
        return {"functions": len(units), "chunks": len(chunks), "source_bytes": len(text.encode()), "all_lines_once": [line for unit in units for line in range(unit["line_start"], unit["line_end"] + 1)] == list(range(1, functions + 1))}

    def pdf_case():
        with fitz.open(output / "physical-three-pages.pdf") as pdf:
            return {"physical_pages": pdf.page_count, "text_pages": [bool(page.get_text().strip()) for page in pdf]}

    cases = [
        ("docx_markers", lambda: docx_case(output / "markers.docx")),
        ("docx_automatic", lambda: docx_case(output / "automatic.docx")),
        ("physical_pdf", pdf_case),
        ("scan_ocr", lambda: {"ocr_text_chars": len(image_text((output / "scan-chi-tra.png").read_bytes()))}),
        ("cpp_400_functions", lambda: cpp_case(source, 400)),
        ("cpp_2000_functions_over_16k_file", lambda: cpp_case(large_source, 2000)),
        ("oversize_atomic_function", lambda: get_chunker("fixed").chunk("int huge() {\n" + " call();\n" * 2200 + "}\n", {"source_units": [{"line_start": 1, "line_end": 2202}]}, {})),
    ]
    measurements = []
    for name, operation in cases:
        tracemalloc.start()
        start = time.perf_counter()
        try:
            value = operation()
            result = {"status": "success", "result": value}
        except Exception as exc:
            result = {"status": "error", "exception_class": type(exc).__name__, "code": getattr(exc, "code", None), "reason": getattr(exc, "details", {}).get("reason")}
        _, peak = tracemalloc.get_traced_memory()
        tracemalloc.stop()
        measurements.append({"case": name, "wall_seconds": round(time.perf_counter()-start, 6), "python_allocation_peak_bytes": peak, "process_peak_rss_kib_cumulative": resource.getrusage(resource.RUSAGE_SELF).ru_maxrss, **result})
    manifest = {path.name: {"bytes": path.stat().st_size, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()} for path in sorted(output.iterdir()) if path.is_file()}
    report = {"scope": "synthetic offline parser/chunker evidence; no provider embedding or production performance claim", "font": "existing PyMuPDF built-in china-t, not a downloaded asset", "memory_limitations": "tracemalloc covers Python allocations; RSS is cumulative process maximum and includes native libraries", "samples": manifest, "measurements": measurements}
    (output / "measurements.safe.json").write_text(json.dumps(report, ensure_ascii=False, indent=2))
    print(json.dumps({"output": str(output), "cases": [{"case": item["case"], "status": item["status"], "code": item.get("code"), "reason": item.get("reason")} for item in measurements]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
